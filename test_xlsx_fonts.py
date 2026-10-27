import os
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import xlsx_fonts as xf  # noqa: E402

FROM, TO, KHF = "Khmer OS Moul Light", "Times New Roman", "Khmer OS System"
KH_TEXT = "ក្រុមហ៊ុន"


def make_book(path):
    from openpyxl import Workbook
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "A"
    ws["A1"] = KH_TEXT + " Trax 2026"
    ws["A1"].font = Font(name=FROM, size=14, bold=True)
    ws["A2"] = "Packing List"
    ws["A2"].font = Font(name=TO, size=12)
    ws["A3"] = KH_TEXT + " Trax 2026"            # same text, other font (shared string)
    ws["A3"].font = Font(name=TO, size=12)
    ws["A4"] = "English only"
    ws["A4"].font = Font(name=FROM, size=11)
    ws["A5"] = CellRichText(TextBlock(InlineFont(rFont=FROM, sz=10), KH_TEXT), " plain")
    ws.merge_cells("B1:C1")
    ws2 = wb.create_sheet("B")
    ws2["A1"] = KH_TEXT
    ws2["A1"].font = Font(name=FROM, size=9)
    wb.save(path)


class Test(unittest.TestCase):
    def test_roundtrip(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as td:
            a, b = os.path.join(td, "a.xlsx"), os.path.join(td, "b.xlsx")
            make_book(a)
            self.assertTrue(xf.needs_font_change(a, FROM))
            st = xf.replace_fonts_xml(a, b, FROM, TO, KHF)
            self.assertGreaterEqual(st["cells"], 3)
            self.assertFalse(xf.needs_font_change(b, FROM))
            zipfile.ZipFile(b).testzip()
            wb = load_workbook(b, rich_text=True)
            ws = wb["A"]
            self.assertEqual(ws["A1"].font.name, TO)
            self.assertEqual(ws["A1"].font.sz, 14)
            self.assertTrue(ws["A1"].font.b)
            self.assertEqual(ws["A4"].font.name, TO)
            self.assertEqual(ws["A2"].font.name, TO)
            # Khmer run uses the Khmer font, Latin run Times New Roman
            fonts = [p.font.rFont for p in ws["A1"].value if hasattr(p, "font")]
            self.assertIn(KHF, fonts)
            self.assertIn(TO, fonts)
            self.assertEqual("".join(str(p) for p in ws["A1"].value), KH_TEXT + " Trax 2026")
            # same text in a Times New Roman cell stays untouched
            self.assertEqual(str(ws["A3"].value), KH_TEXT + " Trax 2026")
            self.assertEqual(ws["A3"].font.name, TO)
            # explicit rich text that named the font
            self.assertIn(KHF, [p.font.rFont for p in ws["A5"].value if hasattr(p, "font")])
            self.assertEqual(str(wb["B"]["A1"].value), KH_TEXT)
            self.assertIn("B1:C1", [str(r) for r in ws.merged_cells.ranges])

    def test_unused_font(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as td:
            a = os.path.join(td, "p.xlsx")
            wb = Workbook()
            wb.active["A1"] = "hello"
            wb.save(a)
            self.assertFalse(xf.needs_font_change(a, FROM))


if __name__ == "__main__":
    unittest.main()
