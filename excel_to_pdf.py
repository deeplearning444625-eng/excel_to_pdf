#!/usr/bin/env python3
"""
Excel -> PDF converter with a font-replacement rule.

Rule (default): every cell/run that uses "Khmer OS Moul Light"
is switched to "Times New Roman" BEFORE the PDF is produced.
Times New Roman has no Khmer glyphs, so Khmer letters inside those cells are
given a real Khmer font (auto-detected, or --khmer-font) while English/numbers
use Times New Roman. This prevents broken Khmer in the PDF.
The original Excel file is never modified.

Usage:
    python excel_to_pdf.py input.xlsx                 # -> input.pdf
    python excel_to_pdf.py input.xlsx -o out/ready.pdf
    python excel_to_pdf.py folder_of_xlsx/ -o pdf_out/   # batch
    python excel_to_pdf.py input.xlsx --from-font "Khmer OS Moul Light" --to-font "Times New Roman"
"""
import argparse
import copy
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont

VERSION = "1.1.0"
FROM_FONT = "Khmer OS Moul Light"
TO_FONT = "Times New Roman"


# Khmer letters + Khmer symbols
KH = r"[\u1780-\u17FF\u19E0-\u19FF]"
KHMER_RUN = re.compile(rf"{KH}+(?:[ \u200b\u200c\u200d]+{KH}+)*")

# Khmer fonts that really contain Khmer glyphs (Times New Roman does NOT)
KHMER_PREFERENCE = [
    "Khmer OS Battambang", "Khmer OS System", "Khmer OS Content", "Khmer OS Siemreap",
    "Khmer UI", "Leelawadee UI", "Noto Sans Khmer", "Noto Serif Khmer", "Khmer OS Muol Light",
]


def installed_fonts() -> set:
    names = set()
    try:  # Windows registry
        import winreg
        for hive, sub in ((winreg.HKEY_LOCAL_MACHINE, r"SOFTWARE\Microsoft\Windows NT\CurrentVersion\Fonts"),
                          (winreg.HKEY_CURRENT_USER, r"SOFTWARE\Microsoft\Windows NT\CurrentVersion\Fonts")):
            try:
                with winreg.OpenKey(hive, sub) as k:
                    i = 0
                    while True:
                        n = winreg.EnumValue(k, i)[0]
                        names.add(re.sub(r"\s*\((TrueType|OpenType)\)$", "", n).lower())
                        i += 1
            except OSError:
                pass
        return names
    except ImportError:
        pass
    try:  # Linux / macOS
        out = subprocess.run(["fc-list", ":", "family"], capture_output=True, text=True).stdout
        for line in out.splitlines():
            for fam in line.split(","):
                names.add(fam.strip().lower())
    except FileNotFoundError:
        pass
    return names


def pick_khmer_font(requested):
    if requested:
        return requested
    have = installed_fonts()
    for f in KHMER_PREFERENCE:
        if any(f.lower() == h or h.startswith(f.lower()) for h in have):
            return f
    print("WARNING: no Khmer font found on this computer. Khmer text may look broken. "
          "Install 'Khmer OS Battambang' / 'Noto Sans Khmer' or pass --khmer-font.")
    return KHMER_PREFERENCE[0]


def split_runs(text, base_inline, khmer_font):
    """Latin/number text stays in the cell font (Times New Roman);
    each Khmer run gets a real Khmer font."""
    parts, pos = [], 0
    for m in KHMER_RUN.finditer(text):
        if m.start() > pos:
            parts.append(text[pos:m.start()])
        f = copy.copy(base_inline)
        f.rFont = khmer_font
        parts.append(TextBlock(f, m.group()))
        pos = m.end()
    if pos < len(text):
        parts.append(text[pos:])
    return parts


def replace_fonts(src: Path, dst: Path, from_font: str, to_font: str, khmer_font: str) -> dict:
    """Copy workbook to dst: from_font -> to_font; Khmer letters get a real Khmer font."""
    wb = load_workbook(src, rich_text=True)
    stats = {"cells": 0, "runs": 0}
    from_l = from_font.strip().lower()

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                is_text = isinstance(v, (str, CellRichText))
                cell_match = bool(cell.font) and (cell.font.name or "").strip().lower() == from_l

                if cell_match:
                    f = copy.copy(cell.font)
                    f.name = to_font
                    cell.font = f
                    stats["cells"] += 1

                if not is_text:
                    continue

                base = InlineFont(rFont=to_font, sz=cell.font.sz, b=cell.font.b, i=cell.font.i,
                                  color=cell.font.color)

                if isinstance(v, str):
                    if cell_match and KHMER_RUN.search(v):
                        parts = split_runs(v, base, khmer_font)
                        cell.value = CellRichText(*parts)
                        stats["runs"] += 1
                else:  # already rich text
                    parts, changed = [], False
                    for part in v:
                        if isinstance(part, TextBlock) and part.font is not None \
                                and (part.font.rFont or "").strip().lower() == from_l:
                            nf = copy.copy(part.font)
                            nf.rFont = to_font
                            parts.extend(split_runs(part.text, nf, khmer_font))
                            changed = True
                            stats["runs"] += 1
                        elif isinstance(part, str) and cell_match:
                            parts.extend(split_runs(part, base, khmer_font))
                            changed = True
                        else:
                            parts.append(part)
                    if changed:
                        cell.value = CellRichText(*parts)

    wb.save(dst)
    return stats


def find_soffice() -> str:
    for name in ("soffice", "libreoffice"):
        p = shutil.which(name)
        if p:
            return p
    for p in (r"C:\Program Files\LibreOffice\program\soffice.exe",
              r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
              "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if Path(p).exists():
            return p
    sys.exit("LibreOffice not found. Either install Microsoft Excel + run: pip install pywin32 "
             "(then use --engine excel), or install LibreOffice from https://www.libreoffice.org")


class ExcelEngine:
    """One hidden Excel instance reused for every file (starting Excel is the slow part)."""

    def __init__(self):
        import win32com.client  # pip install pywin32
        self.app = win32com.client.DispatchEx("Excel.Application")
        self.app.Visible = False
        self.app.DisplayAlerts = False
        self.app.ScreenUpdating = False
        self.app.EnableEvents = False
        self.app.AskToUpdateLinks = False

    def convert(self, xlsx: Path, out_pdf: Path) -> None:
        out_pdf.parent.mkdir(parents=True, exist_ok=True)
        wb = self.app.Workbooks.Open(str(xlsx.resolve()), UpdateLinks=0, ReadOnly=True)
        try:
            wb.ExportAsFixedFormat(0, str(out_pdf.resolve()))  # 0 = PDF
        finally:
            wb.Close(False)

    def close(self):
        try:
            self.app.Quit()
        except Exception:
            pass


def xlsx_to_pdf_libreoffice(xlsx: Path, out_pdf: Path) -> None:
    with tempfile.TemporaryDirectory() as td:
        cmd = [find_soffice(), "--headless", "--convert-to", "pdf",
               "--outdir", td, str(xlsx)]
        r = subprocess.run(cmd, capture_output=True, text=True, timeout=180)
        produced = Path(td) / (xlsx.stem + ".pdf")
        if r.returncode != 0 or not produced.exists():
            sys.exit(f"Conversion failed:\n{r.stdout}\n{r.stderr}")
        out_pdf.parent.mkdir(parents=True, exist_ok=True)
        shutil.move(str(produced), out_pdf)


def has_excel() -> bool:
    if os.name != "nt":
        return False
    try:
        import win32com.client  # noqa: F401
        win32com.client.DispatchEx("Excel.Application").Quit()
        return True
    except Exception:
        return False


ENGINE = "auto"
_excel = None


def xlsx_to_pdf(xlsx: Path, out_pdf: Path) -> None:
    global _excel
    engine = ENGINE
    if engine == "auto":
        engine = "excel" if has_excel() else "libreoffice"
    if engine == "excel":
        if _excel is None:
            _excel = ExcelEngine()
        _excel.convert(xlsx, out_pdf)
    else:
        xlsx_to_pdf_libreoffice(xlsx, out_pdf)


def needs_font_change(src: Path, from_font: str) -> bool:
    """Fast pre-check (reads only styles + strings from the zip). If the font is not
    used, the file is converted as-is: faster, and nothing in it is touched."""
    key = from_font.lower().encode()
    try:
        with zipfile.ZipFile(src) as z:
            for name in z.namelist():
                if name == "xl/styles.xml" or name == "xl/sharedStrings.xml" \
                        or name.startswith("xl/worksheets/sheet"):
                    if key in z.read(name).lower():
                        return True
    except zipfile.BadZipFile:
        return True
    return False


def convert_one(src: Path, out_pdf: Path, from_font: str, to_font: str, khmer_font: str) -> None:
    import time
    t0 = time.time()
    if not needs_font_change(src, from_font):
        xlsx_to_pdf(src, out_pdf)
        note = "font not used - converted directly"
    else:
        with tempfile.TemporaryDirectory() as td:
            tmp = Path(td) / src.name
            stats = replace_fonts(src, tmp, from_font, to_font, khmer_font)
            xlsx_to_pdf(tmp, out_pdf)
        note = f"fonts changed: {stats['cells']} cells, {stats['runs']} Khmer runs"
    print(f"OK  {src.name} -> {out_pdf}  ({note}) {time.time() - t0:.1f}s")


def pick_files():
    """Double-click mode: ask the user which Excel files to convert."""
    import tkinter as tk
    from tkinter import filedialog
    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)
    files = filedialog.askopenfilenames(title="Select Excel files to convert to PDF",
                                        filetypes=[("Excel", "*.xlsx *.xlsm")])
    root.destroy()
    return [Path(f) for f in files]


def main():
    global ENGINE
    ap = argparse.ArgumentParser(description="Excel to PDF with font replacement rule")
    ap.add_argument("input", nargs="*", help=".xlsx file(s) or folder (drag & drop works)")
    ap.add_argument("-o", "--output", help="output PDF file (single) or folder")
    ap.add_argument("--from-font", default=FROM_FONT)
    ap.add_argument("--to-font", default=TO_FONT)
    ap.add_argument("--engine", choices=["auto", "excel", "libreoffice"], default="auto",
                    help="auto = Microsoft Excel if available (Windows), else LibreOffice")
    ap.add_argument("--khmer-font", help="font used for Khmer letters (auto-detected if omitted)")
    ap.add_argument("--version", action="version", version=f"excel2pdf {VERSION}")
    a = ap.parse_args()
    ENGINE = a.engine

    inputs = [Path(p) for p in a.input] or pick_files()
    if not inputs:
        return
    kh = pick_khmer_font(a.khmer_font)
    print(f"Rule: '{a.from_font}' -> '{a.to_font}' | Khmer letters font: '{kh}'")

    jobs = []  # (source, pdf)
    for src in inputs:
        if src.is_dir():
            out_dir = Path(a.output) if a.output else src / "pdf"
            jobs += [(f, out_dir / (f.stem + ".pdf")) for f in sorted(src.glob("*.xls[xm]"))
                     if not f.name.startswith("~$")]
        else:
            if a.output and len(inputs) == 1 and Path(a.output).suffix.lower() == ".pdf":
                out = Path(a.output)
            elif a.output:
                out = Path(a.output) / (src.stem + ".pdf")
            else:
                out = src.with_suffix(".pdf")
            jobs.append((src, out))
    if not jobs:
        sys.exit("No .xlsx files found.")

    try:
        for src, out in jobs:
            convert_one(src, out, a.from_font, a.to_font, kh)
    finally:
        if _excel is not None:
            _excel.close()


def pause():
    if sys.stdin and sys.stdin.isatty():
        input("Press Enter to close...")


if __name__ == "__main__":
    try:
        main()
    except SystemExit as e:
        if e.code not in (None, 0):
            print(e.code)
            pause()
        raise
    except Exception as e:  # show errors when the .exe is double-clicked
        print(f"ERROR: {e}")
        pause()
        raise SystemExit(1)
