"""Direct engine: reads the .xlsx data (openpyxl) and draws the PDF itself (reportlab).

Microsoft Excel / LibreOffice are NOT started, so it is much faster. Trade-offs:
  * layout is an approximation of Excel (no charts, shapes, rotated text, headers/footers)
  * Khmer text cannot be shaped here (no HarfBuzz): files with Khmer letters raise NeedsShaping
    so the caller can use Excel/LibreOffice for them
  * formulas show the values saved in the file (same as the Excel engine with manual calculation)
"""
import datetime as _dt
import io
import math
import os
import re
from pathlib import Path

KHMER = re.compile(r"[\u1780-\u17FF\u19E0-\u19FF]")

PAPER = {1: (612, 792), 5: (612, 1008), 7: (522, 756), 8: (842, 1191), 9: (595, 842),
         11: (420, 595), 13: (516, 729)}  # Letter, Legal, Executive, A3, A4, A5, B5
THEME = ["FFFFFF", "000000", "E7E6E6", "44546A", "4472C4", "ED7D31", "A5A5A5", "FFC000", "5B9BD5", "70AD47"]
BORDER_W = {"hair": 0.25, "thin": 0.5, "dotted": 0.5, "dashed": 0.5, "dashDot": 0.5, "dashDotDot": 0.5,
            "medium": 1.0, "mediumDashed": 1.0, "mediumDashDot": 1.0, "mediumDashDotDot": 1.0,
            "slantDashDot": 1.0, "double": 1.2, "thick": 1.8}


class NeedsShaping(Exception):
    """The workbook has Khmer text that this engine cannot draw correctly."""


# ------------------------------------------------------------------ fonts
_font_cache = {}
_FONT_DIRS = [os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts"),
              "/usr/share/fonts/truetype/liberation", "/usr/share/fonts/truetype/dejavu",
              "/Library/Fonts", "/System/Library/Fonts/Supplemental"]
_TTF = {  # family -> (regular, bold, italic, bold-italic) file names
    "times": ("times.ttf", "timesbd.ttf", "timesi.ttf", "timesbi.ttf"),
    "sans": ("arial.ttf", "arialbd.ttf", "ariali.ttf", "arialbi.ttf"),
    "mono": ("cour.ttf", "courbd.ttf", "couri.ttf", "courbi.ttf"),
}
_LIB = {"times": ("LiberationSerif-Regular.ttf", "LiberationSerif-Bold.ttf", "LiberationSerif-Italic.ttf",
                  "LiberationSerif-BoldItalic.ttf"),
        "sans": ("LiberationSans-Regular.ttf", "LiberationSans-Bold.ttf", "LiberationSans-Italic.ttf",
                 "LiberationSans-BoldItalic.ttf"),
        "mono": ("LiberationMono-Regular.ttf", "LiberationMono-Bold.ttf", "LiberationMono-Italic.ttf",
                 "LiberationMono-BoldItalic.ttf")}
_BUILTIN = {"times": ("Times-Roman", "Times-Bold", "Times-Italic", "Times-BoldItalic"),
            "sans": ("Helvetica", "Helvetica-Bold", "Helvetica-Oblique", "Helvetica-BoldOblique"),
            "mono": ("Courier", "Courier-Bold", "Courier-Oblique", "Courier-BoldOblique")}


def _family(name: str, from_font: str, to_font: str) -> str:
    n = (name or "").strip().lower()
    if n == from_font.strip().lower():
        n = to_font.lower()  # the rule: "Khmer OS Moul Light" -> "Times New Roman"
    if "times" in n or "serif" in n and "sans" not in n or "cambria" in n or "georgia" in n:
        return "times"
    if "courier" in n or "consolas" in n or "mono" in n:
        return "mono"
    return "sans"


def _font(family: str, bold: bool, italic: bool) -> str:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    idx = (1 if bold else 0) + (2 if italic else 0)
    key = (family, idx)
    if key in _font_cache:
        return _font_cache[key]
    chosen = _BUILTIN[family][idx]
    for files in (_TTF[family], _LIB[family]):  # real TrueType = full Unicode coverage
        for d in _FONT_DIRS:
            p = os.path.join(d, files[idx])
            if os.path.exists(p):
                try:
                    nm = f"x2p-{family}-{idx}"
                    pdfmetrics.registerFont(TTFont(nm, p))
                    chosen = nm
                    break
                except Exception:
                    pass
        if chosen.startswith("x2p-"):
            break
    _font_cache[key] = chosen
    return chosen


def _safe(text: str, fontname: str) -> str:
    """Built-in fonts only know Latin-1/WinAnsi; replace what they cannot draw."""
    if fontname.startswith("x2p-"):
        return text
    return text.encode("cp1252", "replace").decode("cp1252")


# ------------------------------------------------------------------ colours
def _tint(hex6: str, tint: float) -> str:
    r, g, b = (int(hex6[i:i + 2], 16) / 255 for i in (0, 2, 4))
    import colorsys
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    l = l * (1 + tint) if tint < 0 else l * (1 - tint) + tint
    r, g, b = colorsys.hls_to_rgb(h, max(0, min(1, l)), s)
    return "%02X%02X%02X" % (round(r * 255), round(g * 255), round(b * 255))


def _rgb(color, default=None):
    """openpyxl Color -> 'RRGGBB' or default."""
    if color is None:
        return default
    try:
        t = color.type
        if t == "rgb" and isinstance(color.rgb, str) and len(color.rgb) in (6, 8):
            return color.rgb[-6:]
        if t == "theme":
            return _tint(THEME[color.theme % 10], color.tint or 0)
        if t == "indexed":
            from openpyxl.styles.colors import COLOR_INDEX
            v = COLOR_INDEX[color.indexed]
            return v[-6:] if isinstance(v, str) and len(v) >= 6 and color.indexed < 64 else default
    except Exception:
        pass
    return default


def _col(hex6):
    from reportlab.lib.colors import HexColor
    return HexColor("#" + hex6)


# ------------------------------------------------------------------ number formats
_MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September",
           "October", "November", "December"]
_DAYS = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]
_TOK = re.compile(r'"[^"]*"|\[[^\]]*\]|\\.|_.|\*.|yyyy|yy|mmmm|mmm|mm|m|dddd|ddd|dd|d|hh|h|ss|s|AM/PM|am/pm|A/P|.',
                  re.I)


def _fmt_date(v, fmt: str) -> str:
    out, last_h = [], False
    toks = _TOK.findall(fmt)
    ampm = any(t.upper() in ("AM/PM", "A/P") for t in toks)
    for i, t in enumerate(toks):
        tl = t.lower()
        if t.startswith('"'):
            out.append(t[1:-1])
        elif t.startswith("[") or t[:1] in "_*":
            continue
        elif t.startswith("\\"):
            out.append(t[1:])
        elif tl == "yyyy":
            out.append("%04d" % v.year)
        elif tl == "yy":
            out.append("%02d" % (v.year % 100))
        elif tl in ("m", "mm"):
            k = i
            before = next((x.lower() for x in reversed(toks[:k]) if x[:1].isalpha()), None)
            after = next((x.lower() for x in toks[k + 1:] if x[:1].isalpha()), None)
            is_min = before in ("h", "hh") or after in ("s", "ss")
            n = v.minute if is_min else v.month
            out.append("%02d" % n if tl == "mm" else str(n))
        elif tl == "mmm":
            out.append(_MONTHS[v.month - 1][:3])
        elif tl == "mmmm":
            out.append(_MONTHS[v.month - 1])
        elif tl == "d":
            out.append(str(v.day))
        elif tl == "dd":
            out.append("%02d" % v.day)
        elif tl == "ddd":
            out.append(_DAYS[v.weekday()][:3])
        elif tl == "dddd":
            out.append(_DAYS[v.weekday()])
        elif tl in ("h", "hh"):
            h = v.hour % 12 or 12 if ampm else v.hour
            out.append("%02d" % h if tl == "hh" else str(h))
        elif tl in ("s", "ss"):
            out.append("%02d" % v.second if tl == "ss" else str(v.second))
        elif t.upper() in ("AM/PM", "A/P"):
            out.append(("PM" if v.hour >= 12 else "AM") if t.upper() == "AM/PM" else ("P" if v.hour >= 12 else "A"))
        else:
            out.append(t)
        last_h = tl in ("h", "hh")
    return "".join(out)


def _fmt_number(v: float, fmt: str):
    secs = fmt.split(";")
    sec, neg_section = secs[0], False
    if v < 0 and len(secs) > 1:
        sec, v, neg_section = secs[1], -v, True
    elif v == 0 and len(secs) > 2:
        sec = secs[2]
    # literal pieces -> placeholders so they cannot be mistaken for digit codes
    lits = []

    def keep(m):
        lits.append(m.group(1) if m.group(1) is not None else m.group(2) if m.group(2) is not None
                    else m.group(3) or "")
        return chr(0xE000 + len(lits) - 1)

    body = re.sub(r'"([^"]*)"|\\(.)|\[\$([^\]\-]*)[^\]]*\]', keep, sec)
    body = re.sub(r"\[[^\]]*\]|_.|\*.", "", body)
    m = re.search(r"[#0?,]*[0#?][#0?,]*(?:\.[0#?]*)?", body)
    pct = "%" in body
    if pct:
        v *= 100
    if not m:  # text only, e.g. "-" for zero
        out = body
    else:
        spec = m.group(0)
        dec = len(spec.split(".")[1]) if "." in spec else 0
        n = f"{abs(v):,.{dec}f}" if "," in spec else f"{abs(v):.{dec}f}"
        intpart = re.sub(r"[^0#?]", "", spec.split(".")[0])
        minw = intpart.count("0")
        ip, _, fp = n.partition(".")
        opt = len(spec.split(".")[1]) - len(spec.split(".")[1].rstrip("#")) if "." in spec else 0
        while opt and fp.endswith("0"):
            fp, opt = fp[:-1], opt - 1
        if "," not in spec and len(ip) < minw:
            ip = ip.rjust(minw, "0")
        if ip == "0" and minw == 0 and intpart:
            ip = ""
        n = ip + ("." + fp if fp else "")
        out = body[:m.start()] + n + body[m.end():]
    out = re.sub("[\ue000-\uf8ff]", lambda k: lits[ord(k.group(0)) - 0xE000], out)
    if v < 0 and not neg_section and len(secs) == 1:
        out = "-" + out
    return out


def format_value(v, fmt: str, is_date: bool) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, (_dt.datetime, _dt.date, _dt.time)):
        if isinstance(v, _dt.date) and not isinstance(v, _dt.datetime):
            v = _dt.datetime(v.year, v.month, v.day)
        if isinstance(v, _dt.time):
            v = _dt.datetime(1900, 1, 1, v.hour, v.minute, v.second)
        f = fmt if fmt and fmt != "General" else "yyyy-mm-dd"
        return _fmt_date(v, f.split(";")[0])
    if isinstance(v, _dt.timedelta):
        t = int(v.total_seconds())
        return "%d:%02d:%02d" % (t // 3600, t % 3600 // 60, t % 60)
    if isinstance(v, (int, float)):
        if not fmt or fmt == "General":
            if isinstance(v, float) and v == int(v) and abs(v) < 1e15:
                return str(int(v))
            return ("%.10g" % v) if isinstance(v, float) else str(v)
        if fmt == "@":
            return str(v)
        try:
            r = _fmt_number(float(v), fmt)
        except Exception:
            r = None
        return r if r is not None else ("%.10g" % v)
    return str(v)


# ------------------------------------------------------------------ layout helpers
def _col_px_to_pt(width_chars) -> float:
    return (width_chars * 7 + 5) * 0.75


def _wrap(text: str, fontname: str, size: float, width: float, stringWidth) -> list:
    lines = []
    for para in text.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        cur = ""
        for word in re.split(r"(?<= )", para):
            if stringWidth(cur + word, fontname, size) <= width or not cur:
                cur += word
            else:
                lines.append(cur.rstrip())
                cur = word
            while stringWidth(cur, fontname, size) > width and len(cur) > 1:  # very long word
                k = len(cur)
                while k > 1 and stringWidth(cur[:k], fontname, size) > width:
                    k -= 1
                lines.append(cur[:k])
                cur = cur[k:]
        lines.append(cur.rstrip())
    return lines


def _has_border(c) -> bool:
    b = c.border
    return any(getattr(b, s) is not None and getattr(b, s).style for s in ("left", "right", "top", "bottom"))


def _print_range(ws):
    """(min_row, min_col, max_row, max_col) of what should be printed."""
    from openpyxl.utils import range_boundaries
    try:
        pa = ws.print_area
        if pa:
            first = (pa if isinstance(pa, str) else pa[0]).split(",")[0].split("!")[-1].replace("$", "")
            c1, r1, c2, r2 = range_boundaries(first)
            return r1, c1, r2, c2
    except Exception:
        pass
    mr = mc = 0
    r0 = c0 = 10 ** 9
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None or (c.fill is not None and c.fill.fill_type == "solid") or _has_border(c):
                mr, mc = max(mr, c.row), max(mc, c.column)
                r0, c0 = min(r0, c.row), min(c0, c.column)
    for img in getattr(ws, "_images", []):
        try:
            a = img.anchor._from
            mr, mc = max(mr, a.row + 1), max(mc, a.col + 1)
            r0, c0 = min(r0, a.row + 1), min(c0, a.col + 1)
        except Exception:
            pass
    if mr == 0:
        return None
    return 1, 1, mr, mc  # start at A1 like Excel when the print area is not set


# ------------------------------------------------------------------ main
def convert(xlsx: Path, out_pdf: Path, cfg: dict, allow_khmer: bool = False) -> str:
    """Write out_pdf. Returns a short note. Raises NeedsShaping if Khmer is found and not allowed."""
    import openpyxl
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfbase.pdfmetrics import stringWidth
    from reportlab.pdfgen import canvas as rl_canvas

    from_font, to_font = cfg.get("from_font", "Khmer OS Moul Light"), cfg.get("to_font", "Times New Roman")
    fit_mode = cfg.get("fit", "auto")
    wb = openpyxl.load_workbook(str(xlsx), data_only=True)
    out_pdf.parent.mkdir(parents=True, exist_ok=True)
    tmp_pdf = out_pdf.with_suffix(".pdf.part")
    cv = None
    notes, skipped_charts, khmer_seen = [], 0, False
    try:
        for ws in wb.worksheets:
            if ws.sheet_state != "visible":
                continue
            rng = _print_range(ws)
            if not rng:
                continue
            r1, c1, r2, c2 = rng

            # ---- grid geometry
            dflt_w = _col_px_to_pt(ws.sheet_format.defaultColWidth or 8.43)
            widths = {c: dflt_w for c in range(c1, c2 + 1)}
            for d in ws.column_dimensions.values():
                if d.min is None:
                    continue
                for c in range(max(d.min, c1), min(d.max or d.min, c2) + 1):
                    widths[c] = 0.0 if d.hidden else _col_px_to_pt(d.width) if d.width else dflt_w
            dflt_h = ws.sheet_format.defaultRowHeight or 15.0
            cells = {}
            for row in ws.iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2):
                for c in row:
                    cells[(c.row, c.column)] = c
            merges = {}
            covered = set()
            for mr in ws.merged_cells.ranges:
                if mr.max_row < r1 or mr.min_row > r2 or mr.max_col < c1 or mr.min_col > c2:
                    continue
                merges[(mr.min_row, mr.min_col)] = (mr.max_row, mr.max_col)
                for rr in range(mr.min_row, mr.max_row + 1):
                    for cc in range(mr.min_col, mr.max_col + 1):
                        if (rr, cc) != (mr.min_row, mr.min_col):
                            covered.add((rr, cc))

            def span_w(ra, ca, rb, cb):
                return sum(widths.get(c, dflt_w) for c in range(ca, cb + 1))

            def style_of(cell):
                f = cell.font
                fam = _family(f.name, from_font, to_font)
                size = float(f.sz or 11)
                return (_font(fam, bool(f.b), bool(f.i)), size, _rgb(f.color, "000000"), bool(f.u), bool(f.strike))

            prepared = {}
            heights = {}
            for r in range(r1, r2 + 1):
                rd = ws.row_dimensions[r] if r in ws.row_dimensions else None
                if rd is not None and rd.hidden:
                    heights[r] = 0.0
                    continue
                fixed = rd.height if rd is not None and rd.height and rd.customHeight is not False else None
                need = dflt_h
                for c in range(c1, c2 + 1):
                    cell = cells.get((r, c))
                    if cell is None or (r, c) in covered or cell.value is None:
                        continue
                    txt = format_value(cell.value, cell.number_format, False)
                    if txt == "":
                        continue
                    if KHMER.search(txt):
                        khmer_seen = True
                        if not allow_khmer:
                            raise NeedsShaping()
                    fname, size, color, ul, st = style_of(cell)
                    al = cell.alignment
                    wrap = bool(al.wrap_text)
                    cw = span_w(r, c, *merges.get((r, c), (r, c))) - 4 - (al.indent or 0) * 6
                    lines = _wrap(txt, fname, size, max(cw, 1), stringWidth) if wrap else txt.split("\n")[:1]
                    prepared[(r, c)] = (txt, lines, fname, size, color, ul, st)
                    if wrap and fixed is None and (r, c) not in merges:
                        need = max(need, len(lines) * size * 1.2 + 3)
                    else:
                        need = max(need, size * 1.25 + 3) if fixed is None else need
                heights[r] = float(fixed) if fixed is not None else need

            tot_w = sum(widths.values())
            tot_h = sum(heights.values())
            if tot_w <= 0 or tot_h <= 0:
                continue

            # ---- page setup
            ps, pm = ws.page_setup, ws.page_margins
            pw, ph = PAPER.get(ps.paperSize or 9, PAPER[9])
            if (ps.orientation or "portrait") == "landscape":
                pw, ph = ph, pw
            ml, mr_, mt, mb = (x * 72 for x in (pm.left, pm.right, pm.top, pm.bottom))
            avail_w, avail_h = pw - ml - mr_, ph - mt - mb
            scale = (ps.scale or 100) / 100.0
            fit_page = bool(ws.sheet_properties.pageSetUpPr and ws.sheet_properties.pageSetUpPr.fitToPage)
            if fit_mode == "all":
                fw, fh = 1, 1
            elif fit_page:
                fw = ps.fitToWidth if ps.fitToWidth is not None else 1
                fh = ps.fitToHeight if ps.fitToHeight is not None else 1
            else:
                fw = fh = None
            if fw is not None:
                scale = 1.0
                if fw:
                    scale = min(scale, avail_w * fw / tot_w)
                if fh:
                    scale = min(scale, avail_h * fh / tot_h)
            scale = max(scale, 0.1)

            # ---- pagination: down, then over (Excel default)
            def groups(sizes, limit):
                out, cur, acc = [], [], 0.0
                for k, s in sizes:
                    if cur and acc + s > limit + 0.01:
                        out.append(cur)
                        cur, acc = [], 0.0
                    cur.append(k)
                    acc += s
                if cur:
                    out.append(cur)
                return out

            col_pages = groups([(c, widths[c]) for c in range(c1, c2 + 1) if widths[c] > 0], avail_w / scale)
            row_pages = groups([(r, heights[r]) for r in range(r1, r2 + 1) if heights[r] > 0], avail_h / scale)
            if cv is None:
                cv = rl_canvas.Canvas(str(tmp_pdf), pagesize=(pw, ph), pageCompression=1)
                cv.setTitle(Path(xlsx).stem)
            else:
                cv.setPageSize((pw, ph))
            center_h = bool(ws.print_options.horizontalCentered)

            for cols in col_pages:
                for rows in row_pages:
                    xs, x = {}, 0.0
                    for c in cols:
                        xs[c] = x
                        x += widths[c]
                    ys, y = {}, 0.0
                    for r in rows:
                        ys[r] = y
                        y += heights[r]
                    area_w, area_h = x, y
                    ox = ml + (max(0, (avail_w - area_w * scale) / 2) if center_h else 0)
                    cv.saveState()
                    cv.translate(ox, ph - mt)
                    cv.scale(scale, scale)
                    p = cv.beginPath()
                    p.rect(0, -area_h, area_w, area_h)
                    cv.clipPath(p, stroke=0, fill=0)

                    def box(r, c):
                        """x, y(top), w, h of the cell/merge starting at (r, c), clipped to this page."""
                        rb, cb = merges.get((r, c), (r, c))
                        ww = sum(widths[k] for k in range(c, cb + 1) if k in xs)
                        hh = sum(heights[k] for k in range(r, rb + 1) if k in ys)
                        return xs[c], ys[r], ww, hh

                    # fills
                    for r in rows:
                        for c in cols:
                            cell = cells.get((r, c))
                            if cell is None or (r, c) in covered:
                                continue
                            f = cell.fill
                            if f is not None and f.fill_type == "solid":
                                hexc = _rgb(f.fgColor)
                                if hexc and hexc.upper() != "FFFFFF":
                                    bx, by, bw, bh = box(r, c)
                                    cv.setFillColor(_col(hexc))
                                    cv.rect(bx, -(by + bh), bw, bh, stroke=0, fill=1)
                    # text
                    for r in rows:
                        for c in cols:
                            pr = prepared.get((r, c))
                            if not pr:
                                continue
                            cell = cells[(r, c)]
                            txt, lines, fname, size, color, ul, st = pr
                            al = cell.alignment
                            bx, by, bw, bh = box(r, c)
                            horiz = al.horizontal or ("right" if isinstance(cell.value, (int, float, _dt.date, _dt.datetime))
                                                      and not isinstance(cell.value, bool) else "left")
                            ind = (al.indent or 0) * 6
                            wrap = bool(al.wrap_text)
                            if al.shrink_to_fit and not wrap:
                                tw = stringWidth(lines[0], fname, size)
                                if tw > bw - 4:
                                    size = max(4, size * (bw - 4) / tw)
                            lh = size * 1.2
                            th = lh * len(lines)
                            va = al.vertical or "bottom"
                            top = by + 1 if va == "top" else by + (bh - th) / 2 if va in ("center", "distributed", "justify") \
                                else by + bh - th - 1
                            # text that does not wrap may spill over empty neighbours
                            cl_x, cl_w = bx, bw
                            if not wrap and (r, c) not in merges:
                                tw = stringWidth(lines[0], fname, size) + 4 + ind
                                if tw > bw:
                                    left = right = 0.0
                                    if horiz in ("left", "general", "fill", "justify", "centerContinuous", "center"):
                                        k = c
                                        while k + 1 in xs and cells.get((r, k + 1)) is not None \
                                                and cells[(r, k + 1)].value is None and right < tw:
                                            k += 1
                                            right += widths[k]
                                    if horiz in ("right", "center"):
                                        k = c
                                        while k - 1 in xs and cells.get((r, k - 1)) is not None \
                                                and cells[(r, k - 1)].value is None and left < tw:
                                            k -= 1
                                            left += widths[k]
                                    cl_x, cl_w = bx - left, bw + left + right
                            cv.saveState()
                            q = cv.beginPath()
                            q.rect(cl_x, -(by + bh), cl_w, bh)
                            cv.clipPath(q, stroke=0, fill=0)
                            cv.setFillColor(_col(color))
                            cv.setStrokeColor(_col(color))
                            cv.setFont(fname, size)
                            for i, line in enumerate(lines):
                                s = _safe(line, fname)
                                w = stringWidth(s, fname, size)
                                if horiz == "right":
                                    tx = cl_x + cl_w - w - 2 - ind
                                elif horiz in ("center", "centerContinuous"):
                                    tx = cl_x + (cl_w - w) / 2
                                else:
                                    tx = bx + 2 + ind
                                ty = -(top + i * lh + size * 1.0)
                                cv.drawString(tx, ty, s)
                                if ul or st:
                                    cv.setLineWidth(max(0.4, size / 24))
                                    if ul:
                                        cv.line(tx, ty - size * 0.12, tx + w, ty - size * 0.12)
                                    if st:
                                        cv.line(tx, ty + size * 0.3, tx + w, ty + size * 0.3)
                            cv.restoreState()
                    # borders (merged cells: use the outer edges of the region)
                    def line(sd, x1, y1, x2, y2):
                        if sd is None or not sd.style:
                            return
                        cv.setStrokeColor(_col(_rgb(sd.color, "000000")))
                        cv.setLineWidth(BORDER_W.get(sd.style, 0.5))
                        if sd.style in ("dashed", "mediumDashed", "dashDot", "dashDotDot"):
                            cv.setDash(3, 2)
                        elif sd.style in ("dotted", "hair"):
                            cv.setDash(1, 1.5)
                        else:
                            cv.setDash()
                        cv.line(x1, -y1, x2, -y2)
                        cv.setDash()

                    for r in rows:
                        for c in cols:
                            if (r, c) in covered:
                                continue
                            cell = cells.get((r, c))
                            if cell is None:
                                continue
                            bx, by, bw, bh = box(r, c)
                            rb, cb = merges.get((r, c), (r, c))
                            b = cell.border
                            right_cell = cells.get((r, cb), cell)
                            bottom_cell = cells.get((rb, c), cell)
                            line(b.top, bx, by, bx + bw, by)
                            line(b.left, bx, by, bx, by + bh)
                            line(right_cell.border.right, bx + bw, by, bx + bw, by + bh)
                            line(bottom_cell.border.bottom, bx, by + bh, bx + bw, by + bh)
                    # pictures that start on this page
                    for img in getattr(ws, "_images", []):
                        try:
                            a = img.anchor._from
                            rr, cc = a.row + 1, a.col + 1
                            if rr not in ys or cc not in xs:
                                continue
                            ix = xs[cc] + (a.colOff or 0) / 12700.0
                            iy = ys[rr] + (a.rowOff or 0) / 12700.0
                            w = (img.width or 0) * 0.75
                            h = (img.height or 0) * 0.75
                            to = getattr(img.anchor, "to", None)
                            if to is not None:  # two-cell anchor: size comes from the end cell
                                w = sum(widths.get(k, dflt_w) for k in range(cc, to.col + 1)) \
                                    + (to.colOff or 0) / 12700.0 - (a.colOff or 0) / 12700.0
                                h = sum(heights.get(k, dflt_h) for k in range(rr, to.row + 1)) \
                                    + (to.rowOff or 0) / 12700.0 - (a.rowOff or 0) / 12700.0
                            data = img._data()
                            cv.drawImage(ImageReader(io.BytesIO(data)), ix, -(iy + h), w, h, mask="auto")
                        except Exception:
                            pass
                    cv.restoreState()
                    cv.showPage()
            skipped_charts += len(getattr(ws, "_charts", []))
        if cv is None:  # nothing to print
            cv = rl_canvas.Canvas(str(tmp_pdf), pagesize=PAPER[9])
            cv.setFont("Helvetica", 10)
            cv.drawString(72, 770, "(empty workbook)")
            cv.showPage()
        cv.save()
        os.replace(tmp_pdf, out_pdf)
    except BaseException:
        try:
            tmp_pdf.unlink()
        except OSError:
            pass
        raise
    finally:
        wb.close()
    if skipped_charts:
        notes.append(f"{skipped_charts} chart(s) not drawn")
    if khmer_seen:
        notes.append("Khmer text drawn without shaping")
    return ", ".join(notes)
